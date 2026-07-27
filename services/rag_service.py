"""
RAG Service v2 - truy xuất knowledge base nội bộ cho AI Test Case Generator.
Tính năng:
- Đọc .md/.txt/.xlsx/.xlsm trong rag/knowledge/.
- Tự dò đường dẫn rag/knowledge dù rag_service.py nằm trong services/ hay cùng cấp project.
- Keyword search + synonym dictionary + domain/chức năng/business-rule boost.
- Có debug log bằng biến môi trường RAG_DEBUG=1.
- Không cần FAISS, chạy nhẹ cho đồ án. Sau này có thể thay _search_documents bằng FAISS.
"""

from __future__ import annotations
import os
import re
import hashlib
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
try:
    import openpyxl
except Exception: 
    openpyxl = None
@dataclass
class RAGDocument:
    source: str
    content: str
    domain: str = ""
    module: str = ""
class RAGService:
    """Keyword/Hybrid RAG đơn giản cho AI Test Case Generator."""  
    SYNONYMS: dict[str, list[str]] = {
        "đăng nhập": ["login", "authentication", "xác thực", "đăng nhâp"],
        "đăng xuất": ["logout", "thoát", "sign out"],
        "thêm mới": ["thêm", "tạo mới", "tạo", "add", "create", "insert", "tự sinh mã", "trường bắt buộc"],
        "cập nhật": ["sửa", "chỉnh sửa", "edit", "update", "modify", "không thay đổi dữ liệu", "xung đột cập nhật"],
        "xóa": ["xoá", "delete", "remove", "xác nhận xóa", "ràng buộc tham chiếu"],
        "hủy bỏ": ["huỷ bỏ", "hủy", "huỷ", "cancel"],
        "sinh mã": ["tạo mã", "generate code", "generate-code"],
        "thêm mới và tiếp tục": ["lưu và tiếp tục", "thêm và tiếp tục", "save and continue"],
        "đóng popup": ["đóng modal", "close popup", "close modal"],
        "tìm kiếm": ["tìm", "search", "filter", "lọc", "tra cứu"],
        "xem chi tiết": ["xem", "view", "detail", "chi tiết"],
        "phân trang": ["pagination", "page", "next", "previous", "trang"],
        "xuất excel": ["export excel", "xlsx", "xuất file", "export", "download"],
        "upload file": ["tải lên", "upload", "đính kèm", "file upload"],
        "quay lại": ["back", "trở về", "return"],
        "phân quyền": ["role", "permission", "authorization", "admin", "hr", "user"],
        "đăng tin việc làm": ["tin tuyển dụng", "job posting", "job post", "đăng tuyển"],
        "ứng tuyển": ["apply", "nộp hồ sơ", "apply job", "cv"],
        "hồ sơ ứng viên": ["candidate profile", "cv", "resume", "ứng viên"],
        "lịch phỏng vấn": ["interview", "schedule", "phỏng vấn"],
        "chuyển khoản": ["transfer", "giao dịch chuyển tiền"],
        "lịch khám": ["appointment", "đặt lịch khám"],
        "đơn thuốc": ["prescription", "kê đơn"],
        "thanh toán": ["payment", "pay", "hóa đơn", "hoá đơn"],
        "nhập điểm": ["điểm", "grade", "score"],
        "lớp học": ["quản lý lớp học", "class", "classroom"],
        "môn học": ["quản lý môn học", "subject", "course"],
        "quản lý tài khoản": ["account management", "thông tin tài khoản"],
        "lịch sử giao dịch": ["transaction history", "lịch sử giao dịch ngân hàng"],
        "báo cáo": ["report", "xuất báo cáo", "thống kê"],
    }
    DOMAIN_KEYWORDS: dict[str, list[str]] = {
        "common": ["common", "dùng chung", "chung"],
        "school": ["trường học", "sinh viên", "giáo viên", "lớp học", "môn học", "điểm", "thời khóa biểu"],
        "hospital": ["bệnh viện", "bệnh nhân", "bác sĩ", "lịch khám", "đơn thuốc", "viện phí"],
        "bank": ["ngân hàng", "tài khoản", "chuyển khoản", "số dư", "giao dịch"],
        "inventory": ["quản lý kho", "kho hàng", "tồn kho", "nhập kho", "xuất kho", "kiểm kê", "nhà cung cấp", "phiếu nhập", "phiếu xuất", "vật tư", "kho", "danh mục kho", "mã kho", "tên kho"],
        "hr": ["nhân sự", "nhân viên", "phòng ban", "chấm công", "lương", "nghỉ phép"],
        "recruitment": ["tuyển dụng", "hr", "ứng viên", "việc làm", "đăng tin", "phỏng vấn", "cv", "hồ sơ"],
        "ecommerce": ["thương mại điện tử", "sản phẩm", "giỏ hàng", "đơn hàng", "checkout"],
        "hotel": ["khách sạn", "đặt phòng", "nhận phòng", "trả phòng"],
        "library": ["thư viện", "sách", "độc giả", "mượn sách", "trả sách"],
    }
    DOMAIN_ALIASES: dict[str, str] = {
        "inventory": "inventory",
        "kho": "inventory",
        "trường": "school",
        "benhvien": "hospital",
        "nganhang": "bank",
        "tuyendung": "recruitment",
        "nhansu": "hr",
        "thương mại điện tử" : "ecommerce",
        "khách sạn" : "hotel",
    }
    IMPORTANT_PHRASES = [
        "đăng nhập", "đăng xuất", "tìm kiếm", "thêm mới", "cập nhật", "chỉnh sửa",
        "xóa", "xoá", "xem chi tiết", "phân trang", "xuất excel", "xuất file",
        "quay lại", "phân quyền", "upload file",
        "hủy bỏ", "huỷ bỏ", "sinh mã", "đóng popup", "thêm mới và tiếp tục",
        "quản lý kho", "quản lý sinh viên", "quản lý bệnh nhân", "quản lý tuyển dụng",
        "đăng tin việc làm", "ứng tuyển", "quản lý hồ sơ", "hồ sơ ứng viên", "lịch phỏng vấn",
        "chuyển khoản", "lịch khám", "đơn thuốc", "thanh toán", "nhập điểm",
        "lớp học", "môn học", "quản lý tài khoản", "lịch sử giao dịch", "báo cáo",
        "common", "school", "hospital", "bank", "inventory", "hr", "recruitment", "ecommerce", "hotel","library"
    ]
    _REQUEST_VERB_META_RE = re.compile(
        r"\b(?:thêm|tạo|bổ\s+sung)\s+(?:thêm\s+)?"
        r"(?:chức\s+năng|test\s*case|module|yêu\s+cầu)\s*",
        re.IGNORECASE,
    )

    def __init__(self, knowledge_dir: str | None = None, max_chars_per_doc: int = 1800):
        self.max_chars_per_doc = max_chars_per_doc
        self.documents: list[RAGDocument] = []
        self.debug = os.environ.get("RAG_DEBUG", "0").lower() in {"1", "true", "yes", "on"}
        env_dir = os.environ.get("RAG_KNOWLEDGE_DIR", "")
        if knowledge_dir:
            self.knowledge_dir = Path(knowledge_dir)
        elif env_dir:
            self.knowledge_dir = Path(env_dir)
        else:
            self.knowledge_dir = self._resolve_default_knowledge_dir()
        self._load_documents()
    def reload(self) -> None:
        """Nạp lại knowledge sau khi thêm hoặc chỉnh sửa file mà không cần tạo service mới."""
        self.documents.clear()
        self._load_documents()
    def _normalize_domain(self, domain: str) -> str:
        value = (domain or "").strip().lower()
        return self.DOMAIN_ALIASES.get(value, value)
    def _resolve_default_knowledge_dir(self) -> Path:
        here = Path(__file__).resolve()
        candidates = [
            here.parent / "rag" / "knowledge",              
            here.parent.parent / "rag" / "knowledge",       
            here.parent.parent.parent / "rag" / "knowledge",
            Path.cwd() / "rag" / "knowledge",               
        ]
        for c in candidates:
            if c.exists():
                return c
        print(
            "[RAG] Không tìm thấy 'rag/knowledge' ở các vị trí đã thử:\n"
            + "\n".join(f"  - {c}" for c in candidates)
        )
        return candidates[0]
    def _load_documents(self) -> None:
        if not self.knowledge_dir.exists():
            print(f"[RAG] Chưa có thư mục knowledge: {self.knowledge_dir}")
            return
        self._seen_content_hashes: set[str] = set()
        self._duplicates_skipped_count = 0
        for path in sorted(self.knowledge_dir.rglob("*")):
            if not path.is_file():
                continue
            suffix = path.suffix.lower()
            try:
                if suffix in {".md", ".txt"}:
                    self._load_text_file(path)
                elif suffix in {".xlsx", ".xlsm"}:
                    self._load_excel_file(path)
            except Exception as exc:
                print(f"[RAG] Bỏ qua file {path.name}: {exc}")
        print(
            f"[RAG] Loaded {len(self.documents)} documents from {self.knowledge_dir}"
            + (
                f" (đã loại {self._duplicates_skipped_count} chunk trùng nội dung)"
                if self._duplicates_skipped_count else ""
            )
        )
    def _is_duplicate_content(self, content: str) -> bool:
        """
        Chuẩn hóa nội dung (bỏ khoảng trắng thừa, lowercase) rồi hash để
        phát hiện chunk trùng lặp GẦN NHƯ HOÀN TOÀN — tránh RAG context bị
        nhồi nhiều bản sao giống hệt nhau (vd cùng nội dung xuất hiện ở cả
        .md và .xlsx, hoặc bị copy-paste giữa nhiều sheet/file domain).
        """
        if not hasattr(self, '_seen_content_hashes'):
            self._seen_content_hashes = set()
        if not hasattr(self, '_duplicates_skipped_count'):
            self._duplicates_skipped_count = 0

        normalized = re.sub(r'\s+', ' ', content.strip().lower())
        content_hash = hashlib.md5(normalized.encode('utf-8')).hexdigest()
        if content_hash in self._seen_content_hashes:
            self._duplicates_skipped_count += 1
            return True
        self._seen_content_hashes.add(content_hash)
        return False

    def _load_text_file(self, path: Path) -> None:
        text = path.read_text(encoding="utf-8", errors="ignore")
        chunks = re.split(r"\n(?=##+\s+|DOMAIN\s*:|MODULE\s*:)", text)
        for idx, chunk in enumerate(chunks, 1):
            chunk = chunk.strip()
            if len(chunk) < 40:
                continue
            domain = self._extract_meta(chunk, "domain") or self._infer_domain_from_path_or_text(path, chunk)
            chức_năng = self._extract_meta(chunk, "chức năng") or self._extract_heading(chunk)
            content = chunk[: self.max_chars_per_doc]
            if self._is_duplicate_content(content):
                continue
            self.documents.append(
                RAGDocument(
                    source=f"{path.relative_to(self.knowledge_dir).as_posix()}#{idx}",
                    content=content,
                    domain=domain,
                    module=chức_năng,
                )
            )

    def _load_excel_file(self, path: Path) -> None:
        if openpyxl is None:
            print("[RAG] openpyxl chưa được cài, bỏ qua Excel.")
            return
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(v).strip() if v is not None else "" for v in rows[0]]
            header_map = {h.lower(): i for i, h in enumerate(headers) if h}

            rag_col = next((i for h, i in header_map.items() if "rag" in h and "document" in h), None)
            domain_col = next((i for h, i in header_map.items() if h in {"domain", "lĩnh vực", "linh vuc"}), None)
            module_col = next((i for h, i in header_map.items() if h in {"module", "chức năng", "chuc nang", "feature"}), None)

            for r_idx, row in enumerate(rows[1:], 2):
                if rag_col is not None:
                    val = row[rag_col] if rag_col < len(row) else None
                    content = str(val).strip() if val else ""
                else:
                    pairs = []
                    for c_idx, h in enumerate(headers):
                        if not h or c_idx >= len(row) or row[c_idx] is None:
                            continue
                        value = str(row[c_idx]).strip()
                        if value:
                            pairs.append(f"{h}: {value}")
                    content = "\n".join(pairs)

                if len(content) < 60:
                    continue
                domain = ""
                module = ""
                if domain_col is not None and domain_col < len(row) and row[domain_col]:
                    domain = str(row[domain_col]).strip().lower()
                if module_col is not None and module_col < len(row) and row[module_col]:
                    module = str(row[module_col]).strip().lower()
                domain = domain or self._infer_domain_from_path_or_text(path, content)
                module = module or self._extract_meta(content, "module")
                final_content = content[: self.max_chars_per_doc]
                if self._is_duplicate_content(final_content):
                    continue
                self.documents.append(
                    RAGDocument(
                        source=f"{path.relative_to(self.knowledge_dir).as_posix()}:{ws.title}:R{r_idx}",
                        content=final_content,
                        domain=domain,
                        module=module,
                    )
                )

    def retrieve(self, query: str, top_k: int = 5, domain: str = "", targeted: bool = False) -> str:
        """
        targeted=True (TH1 — user chỉ định cụ thể chức năng cần sinh):
        sau khi dò được các module thực sự được yêu cầu, LỌC CỨNG các
        chunk không thuộc module đó ra khỏi context cuối cùng đưa vào
        prompt AI (xem _search_documents) — không để RAG tự bổ sung
        module ngoài phạm vi yêu cầu.
        targeted=False (mặc định, TH2 — full scan / không xác định rõ
        phạm vi): giữ nguyên hành vi cũ, không lọc cứng theo module.
        """
        if not self.documents:
            return ""
        ranked = self._search_documents(query, top_k=top_k, domain=domain, targeted=targeted)
        if not ranked:
            print(f"[RAG] Domain: {domain or '(auto)'}")
            print("[RAG] Requested modules: (không xác định)")
            print("[RAG] Retrieved modules: (không xác định)")
            print("[RAG] Applied modules: (không xác định)")
            print("[RAG] Retrieved: 0 chunks")
            return ""

        if self.debug:
            print("\n========== RAG DEBUG ==========")
            print("QUERY:")
            print(query[:1000])
            print("\nRESULTS:")
            for score, doc in ranked:
                print(f"- score={score} source={doc.source} domain={doc.domain} module={doc.module}")
            print("===============================\n")

        blocks = []
        for rank, (score, doc) in enumerate(ranked, 1):
            meta = []
            if doc.domain:
                meta.append(f"domain={doc.domain}")
            if doc.module:
                meta.append(f"module={doc.module}")
            meta_text = " ".join(meta)
            blocks.append(
                f"[RAG-{rank}] source={doc.source} score={score} {meta_text}\n"
                f"{doc.content}"
            )
        return "\n\n---\n\n".join(blocks)

    def _search_documents(
        self, query: str, top_k: int = 5, domain: str = "", targeted: bool = False,
    ) -> list[tuple[int, RAGDocument]]:
        query_terms = self._tokenize(query)
        if not query_terms:
            return []

        q_lower = (query or "").lower()
        explicit_domain = self._normalize_domain(domain)
        primary_domain = explicit_domain or self._select_primary_domain(q_lower)
        query_domains = {primary_domain} if primary_domain else self._detect_domains(q_lower)
        query_modules = self._detect_modules(q_lower)
        query_phrases = list(self._important_phrases(q_lower))
        domain_lower = primary_domain
        candidate_docs = self.documents
        if domain_lower:
            candidate_docs = [
                d for d in self.documents
                if not d.domain or self._normalize_domain(d.domain) in {domain_lower, "common"}
            ]

        scored: list[tuple[int, RAGDocument]] = []
        for doc in candidate_docs:
            d_lower = doc.content.lower()
            doc_terms = self._tokenize(doc.content)
            if not doc_terms:
                continue
            score = 0
            common_terms = query_terms & doc_terms
            score += len(common_terms) * 2
            for phrase in query_phrases:
                if phrase in d_lower:
                    score += 12
            for canonical, variants in self.SYNONYMS.items():
                if self._contains_any(q_lower, [canonical, *variants]) and self._contains_any(d_lower, [canonical, *variants]):
                    score += 10
            doc_domain = (doc.domain or "").lower()
            if doc_domain == "common":
                score += 3
            if query_domains:
                if doc_domain in query_domains:
                    score += 25
                elif doc_domain and doc_domain not in query_domains and doc_domain != "common":
                    score -= 6
            doc_module = (doc.module or self._extract_meta(d_lower, "module") or "").lower()
            if doc_module:
                for module in query_modules:
                    if module in doc_module or doc_module in module:
                        score += 22
                if doc_module and doc_module in q_lower:
                    score += 18
            if any(marker in d_lower for marker in ["business_rules", "business rules", "business rule", "quy tắc nghiệp vụ"]):
                score += 8
            if any(marker in d_lower for marker in ["workflow", "luồng nghiệp vụ", "feature", "description"]):
                score += 5
            source_lower = doc.source.lower()
            for domain in query_domains:
                if domain in source_lower:
                    score += 8

            if score > 0:
                scored.append((score, doc))

        scored.sort(key=lambda x: x[0], reverse=True)
        if not scored:
            return []
        MIN_RELEVANCE_SCORE = 8
        RELATIVE_RATIO = 0.35
        top_score = scored[0][0]
        relevance_floor = max(MIN_RELEVANCE_SCORE, top_score * RELATIVE_RATIO)
        relevant = [(s, d) for s, d in scored if s >= relevance_floor][: max(1, top_k)]

        def _chunk_module_names(pairs: list[tuple[int, RAGDocument]]) -> list[str]:
            names: list[str] = []
            for _, d in pairs:
                m = (d.module or "").strip().lower()
                if m and m not in names:
                    names.append(m)
            return names
        applied = relevant
        if targeted and query_modules:
            def _keep(doc: RAGDocument) -> bool:
                doc_module = (doc.module or "").strip().lower()
                if not doc_module:
                    return True
                return any(m in doc_module or doc_module in m for m in query_modules)
            filtered = [(s, d) for s, d in relevant if _keep(d)]
            if filtered:
                applied = filtered

        used_domain = domain_lower or (next(iter(query_domains)) if len(query_domains) == 1 else "")
        requested_modules = sorted(query_modules) if query_modules else []
        retrieved_modules = _chunk_module_names(relevant)
        applied_modules = _chunk_module_names(applied)
        print(f"[RAG] Domain: {used_domain or '(auto)'}")
        print(f"[RAG] Requested modules: {', '.join(requested_modules) if requested_modules else '(auto)'}")
        print(f"[RAG] Retrieved modules: {', '.join(retrieved_modules) if retrieved_modules else '(auto)'}")
        print(f"[RAG] Applied modules: {', '.join(applied_modules) if applied_modules else '(auto)'}")
        print(f"[RAG] Retrieved: {len(applied)} chunks")

        return applied

    def _tokenize(self, text: str) -> set[str]:
        text = (text or "").lower()
        words = re.findall(r"[a-zA-ZÀ-ỹ0-9_]+", text)
        stopwords = {
            "và", "hoặc", "cho", "của", "theo", "một", "các", "trong", "ngoài",
            "người", "dùng", "hệ", "thống", "mô", "tả", "ui", "test", "case",
            "module", "chức", "năng", "quản", "lý", "với", "role", "ứng", "dụng",
        }
        tokens = {w for w in words if len(w) >= 2 and w not in stopwords}
        for canonical, variants in self.SYNONYMS.items():
            if self._contains_any(text, [canonical, *variants]):
                tokens.add(canonical)
                tokens.update(v for v in variants if len(v) >= 2)
        return tokens

    def _important_phrases(self, query_lower: str) -> Iterable[str]:
        return [p for p in self.IMPORTANT_PHRASES if p in query_lower]

    def _contains_term(self, text: str, term: str) -> bool:
        """Khớp theo biên từ để tránh 'kho' khớp nhầm trong 'khoản'."""
        term = (term or "").strip().lower()
        if not term:
            return False
        pattern = rf"(?<![0-9A-Za-zÀ-ỹ_]){re.escape(term)}(?![0-9A-Za-zÀ-ỹ_])"
        return re.search(pattern, (text or "").lower()) is not None

    def _contains_any(self, text: str, terms: list[str]) -> bool:
        return any(self._contains_term(text, t) for t in terms if t)

    def _domain_match_counts(self, text_lower: str) -> dict[str, int]:
        counts: dict[str, int] = {}
        for domain, keywords in self.DOMAIN_KEYWORDS.items():
            hits = sum(1 for keyword in keywords if self._contains_term(text_lower, keyword))
            if hits:
                counts[domain] = hits
        return counts
    def _select_primary_domain(self, text_lower: str) -> str:
        """Chọn một domain chính để chỉ nạp common + domain đó vào prompt."""
        counts = self._domain_match_counts(text_lower)
        counts.pop("common", None)
        if not counts:
            return ""
        max_hits = max(counts.values())
        candidates = {domain for domain, hits in counts.items() if hits == max_hits}
        priority = ["recruitment", "inventory", "school", "hospital", "bank", "hr",
                    "ecommerce", "hotel", "library"]
        for domain in priority:
            if domain in candidates:
                return domain
        return sorted(candidates)[0]

    def _detect_domains(self, text_lower: str) -> set[str]:
        return set(self._domain_match_counts(text_lower))

    def _strip_request_verbs(self, text_lower: str) -> str:
        """
        Bỏ cụm "động từ yêu cầu + từ điều khiển" ('thêm chức năng',
        'tạo testcase', 'bổ sung (thêm) chức năng', ...) TRƯỚC KHI dò
        module, để "thêm"/"tạo" không còn đứng riêng lẻ trong text và bị
        _contains_any() khớp nhầm vào biến thể của canonical "thêm mới".
        Chỉ xóa đúng phần ĐỘNG TỪ + TỪ ĐIỀU KHIỂN, giữ nguyên phần tên
        chức năng phía sau — nên nếu tên chức năng thật sự là "thêm mới"
        (vd "tạo chức năng thêm mới") thì cụm "thêm mới" còn lại vẫn được
        dò bình thường ở bước sau, không bị mất.
        Ví dụ:
          "thêm chức năng đăng xuất"        -> " đăng xuất"
          "thêm testcase đăng nhập"         -> " đăng nhập"
          "tạo chức năng thêm mới"          -> "thêm mới"
          "bổ sung thêm chức năng tìm kiếm" -> "tìm kiếm"
        """
        return self._REQUEST_VERB_META_RE.sub(' ', text_lower)
    def _detect_modules(self, text_lower: str) -> set[str]:
        cleaned = self._strip_request_verbs(text_lower)
        modules = set()
        for canonical, variants in self.SYNONYMS.items():
            if self._contains_any(cleaned, [canonical, *variants]):
                modules.add(canonical)
        for phrase in self.IMPORTANT_PHRASES:
            if phrase in cleaned:
                modules.add(phrase)
        return modules
    def _extract_meta(self, text: str, key: str) -> str:
        m = re.search(rf"^\s*\*?\*?{re.escape(key)}\*?\*?\s*:\s*(.+)$", text, flags=re.I | re.M)
        if not m:
            return ""
        return m.group(1).strip().lower()
    def _extract_heading(self, text: str) -> str:
        m = re.search(r"^#{1,6}\s+(.+)$", text, flags=re.M)
        return m.group(1).strip().lower() if m else ""
    def _infer_domain_from_path_or_text(self, path: Path, text: str) -> str:
        """
        Suy ra domain của 1 đoạn tài liệu.
        Trước đây hàm này chỉ lấy domain ĐẦU TIÊN khớp theo 1 thứ tự ưu
        tiên CỐ ĐỊNH, nên nếu domain đứng trước trong thứ tự (vd
        "inventory") chỉ khớp NHẦM 1 từ khóa ngắn/dễ đụng hàng (vd "kho"
        là substring của "khoản" trong "tài khoản"/"chuyển khoản") thì vẫn
        thắng domain THỰC SỰ đúng (vd "bank") dù domain đó khớp nhiều từ
        khóa hơn hẳn. Sửa: đếm SỐ từ khóa khớp của mỗi domain, domain nào
        khớp NHIỀU nhất thắng; chỉ dùng thứ tự ưu tiên cố định để phá hòa.
        Ngoài ra cộng thêm điểm bonus nếu đường dẫn file nằm trong thư mục
        đặt đúng tên domain (vd "rag/knowledge/bank/..."), vì đây là tín
        hiệu rất đáng tin cậy khi knowledge base được tổ chức theo domain.
        """
        combined = f"{path.as_posix().lower()}\n{text.lower()}"
        path_lower = path.as_posix().lower()
        counts: dict[str, int] = {}
        for domain, keywords in self.DOMAIN_KEYWORDS.items():
            hits = sum(1 for k in keywords if self._contains_term(combined, k))
            if re.search(rf"(^|/){re.escape(domain)}(/|-|_|\.)", path_lower):
                hits += 3
            if hits:
                counts[domain] = hits
        if not counts:
            return ""
        max_hits = max(counts.values())
        candidates = {d for d, h in counts.items() if h == max_hits}
        for d in ["recruitment", "inventory", "school", "hospital", "bank", "hr", "ecommerce", "hotel", "library", "common"]:
            if d in candidates:
                return d
        return next(iter(candidates))